"""
DXF 转 Excel 服务
上传 DXF 文件，解析后输出 Excel（包含 Layers / Lines / Circles / Arcs / Polyline / Rectangles / Annotations）
"""
from __future__ import annotations

import io
import math
import logging
import tempfile
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

LOG = logging.getLogger("dxf_to_excel")

router = APIRouter(prefix="/dxf", tags=["DXF to Excel"])

WORK_DIR = Path(tempfile.gettempdir()) / "dxf_to_excel"
WORK_DIR.mkdir(exist_ok=True)


def _read_dxf(dxf_path: Path):
    """读取 DXF 文件，自动尝试多种编码"""
    import ezdxf

    doc = None
    for encoding in ["utf-8", "gbk", "gb2312", "gb18030", "cp936"]:
        try:
            doc = ezdxf.readfile(str(dxf_path), encoding=encoding)
            break
        except (UnicodeDecodeError, ezdxf.DXFError):
            continue

    if doc is None:
        doc = ezdxf.readfile(str(dxf_path), encoding="utf-8", errors="replace")

    return doc


def _is_rectangle(points: list[tuple[float, float]], is_closed: bool) -> bool:
    """判断多段线是否为矩形（4 个顶点 + 封闭，且所有角都是 90 度）"""
    # 封闭的 LWPOLYLINE 最后一点可能与第一点重合
    pts = list(points)
    if len(pts) == 5 and _pts_close(pts[0], pts[4]):
        pts = pts[:4]
    if len(pts) != 4:
        return False
    if not is_closed and not _pts_close(pts[0], pts[-1]):
        return False

    # 检查 4 条边是否两两垂直（点积为 0）
    for i in range(4):
        p1 = pts[i]
        p2 = pts[(i + 1) % 4]
        p3 = pts[(i + 2) % 4]
        dx1, dy1 = p2[0] - p1[0], p2[1] - p1[1]
        dx2, dy2 = p3[0] - p2[0], p3[1] - p2[1]
        dot = dx1 * dx2 + dy1 * dy2
        if abs(dot) > 1e-4:
            return False
    return True


def _pts_close(a: tuple[float, float], b: tuple[float, float]) -> bool:
    return abs(a[0] - b[0]) < 1e-6 and abs(a[1] - b[1]) < 1e-6


def _rect_center_wh(points: list[tuple[float, float]]) -> tuple[tuple[float, float], float, float]:
    """计算矩形中心点和宽高"""
    xs = [p[0] for p in points[:4]]
    ys = [p[1] for p in points[:4]]
    min_x, max_x = min(xs), max(xs)
    min_y, max_y = min(ys), max(ys)
    cx = round((min_x + max_x) / 2, 4)
    cy = round((min_y + max_y) / 2, 4)
    w = round(max_x - min_x, 4)
    h = round(max_y - min_y, 4)
    return (cx, cy), w, h


def _clean_text(text: str) -> str:
    """清理文本，去除 XML 非法字符"""
    if not text:
        return ""
    text = text.encode("utf-8", errors="replace").decode("utf-8", errors="replace")
    # 移除 XML 1.0 非法字符 (xlsx 基于 XML)
    return "".join(
        ch for ch in text
        if ch == "\t" or ch == "\n" or ch == "\r"
        or ("\x20" <= ch <= "\ud7ff")
        or ("\ue000" <= ch <= "\ufffd")
    )


def dxf_to_excel_bytes(dxf_path: Path) -> bytes:
    """解析 DXF 文件并生成 Excel bytes"""
    from openpyxl import Workbook

    doc = _read_dxf(dxf_path)
    msp = doc.modelspace()

    wb = Workbook()

    # ===== Sheet: Layers =====
    ws_layers = wb.active
    ws_layers.title = "Layers"
    ws_layers.append(["Layer_Name", "Color", "Linetype", "Lineweight", "IsOn", "IsLocked"])
    for layer in doc.layers:
        ws_layers.append([
            _clean_text(layer.dxf.name),
            layer.dxf.color,
            _clean_text(layer.dxf.linetype),
            layer.dxf.lineweight,
            layer.is_on(),
            layer.is_locked(),
        ])

    # ===== Collect entities =====
    lines_data = []
    circles_data = []
    arcs_data = []
    polylines_data = []
    rectangles_data = []
    annotations_data = []

    line_counter = 0
    circle_counter = 0
    arc_counter = 0
    poly_counter = 0
    anno_counter = 0

    for entity in msp:
        etype = entity.dxftype()
        layer_name = _clean_text(entity.dxf.layer)

        if etype == "LINE":
            line_counter += 1
            start = entity.dxf.start
            end = entity.dxf.end
            linetype = _clean_text(entity.dxf.get("linetype", "Continuous"))
            lines_data.append([
                f"LINE_{line_counter}",
                layer_name,
                round(start.x, 4),
                round(start.y, 4),
                round(end.x, 4),
                round(end.y, 4),
                linetype,
            ])

        elif etype == "CIRCLE":
            circle_counter += 1
            center = entity.dxf.center
            circles_data.append([
                f"CIRCLE_{circle_counter}",
                layer_name,
                round(center.x, 4),
                round(center.y, 4),
                round(entity.dxf.radius, 4),
            ])

        elif etype == "ARC":
            arc_counter += 1
            center = entity.dxf.center
            start_angle = entity.dxf.start_angle
            end_angle = entity.dxf.end_angle
            total_angle = (end_angle - start_angle) % 360
            arcs_data.append([
                f"ARC_{arc_counter}",
                layer_name,
                round(center.x, 4),
                round(center.y, 4),
                round(entity.dxf.radius, 4),
                round(start_angle, 4),
                round(total_angle, 4),
            ])

        elif etype == "LWPOLYLINE":
            poly_counter += 1
            points = [(round(p[0], 4), round(p[1], 4)) for p in entity.get_points(format="xy")]
            is_closed = entity.closed

            if _is_rectangle(points, is_closed):
                center, w, h = _rect_center_wh(points)
                rectangles_data.append([
                    f"LWPOLYLINE_{poly_counter}",
                    layer_name,
                    f"({center[0]}, {center[1]})",
                    w,
                    h,
                ])
            else:
                polylines_data.append([
                    f"LWPOLYLINE_{poly_counter}",
                    layer_name,
                    str(points),
                    "是" if is_closed else "否",
                ])

        elif etype == "POLYLINE":
            poly_counter += 1
            points = [
                (round(v.dxf.location.x, 4), round(v.dxf.location.y, 4))
                for v in entity.vertices
            ]
            is_closed = entity.is_closed

            if _is_rectangle(points, is_closed):
                center, w, h = _rect_center_wh(points)
                rectangles_data.append([
                    f"POLYLINE_{poly_counter}",
                    layer_name,
                    f"({center[0]}, {center[1]})",
                    w,
                    h,
                ])
            else:
                polylines_data.append([
                    f"POLYLINE_{poly_counter}",
                    layer_name,
                    str(points),
                    "是" if is_closed else "否",
                ])

        elif etype in ("TEXT", "MTEXT"):
            anno_counter += 1
            if etype == "MTEXT":
                insert = entity.dxf.insert
                text = _clean_text(entity.text)
                height = entity.dxf.get("char_height", entity.dxf.get("height", 0))
            else:
                insert = entity.dxf.insert
                text = _clean_text(entity.dxf.text)
                height = entity.dxf.get("height", 0)

            color = entity.dxf.get("color", 256)  # 256 = BYLAYER
            style = _clean_text(entity.dxf.get("style", ""))
            annotations_data.append([
                f"标注_{anno_counter}",
                text,
                0,
                round(height, 4),
                color,
                round(insert.x, 4),
                round(insert.y, 4),
                layer_name,
                style,
            ])

    # ===== Sheet: Lines =====
    ws_lines = wb.create_sheet("Lines")
    ws_lines.append(["ID", "Layer_Name", "Start_X", "Start_Y", "End_X", "End_Y", "Linetype"])
    for row in lines_data:
        ws_lines.append(row)

    # ===== Sheet: Circles =====
    ws_circles = wb.create_sheet("Circles")
    ws_circles.append(["ID", "Layer_Name", "Center_X", "Center_Y", "Radius"])
    for row in circles_data:
        ws_circles.append(row)

    # ===== Sheet: Arcs =====
    ws_arcs = wb.create_sheet("Arcs")
    ws_arcs.append(["ID", "Layer_Name", "Center_X", "Center_Y", "Radius", "Start_Angle", "Total_Angle"])
    for row in arcs_data:
        ws_arcs.append(row)

    # ===== Sheet: Polyline =====
    ws_poly = wb.create_sheet("Polyline")
    ws_poly.append(["ID", "Layer_Name", "顶点坐标列表", "是否封闭"])
    for row in polylines_data:
        ws_poly.append(row)

    # ===== Sheet: Rectangles =====
    ws_rect = wb.create_sheet("Rectangles")
    ws_rect.append(["ID", "Layer_Name", "中心坐标", "宽", "高"])
    for row in rectangles_data:
        ws_rect.append(row)

    # ===== Sheet: Annotations =====
    ws_anno = wb.create_sheet("Annotations")
    ws_anno.append(["ID", "Annotation_Name", "Distance", "Font_Size", "Font_Color",
                    "Position_X", "Position_Y", "Layer_Name", "Text_Style"])
    for row in annotations_data:
        ws_anno.append(row)

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@router.post("/to-excel")
async def dxf_to_excel_endpoint(
    file: UploadFile = File(..., description="DXF 文件"),
):
    """
    上传 DXF 文件，解析后返回 Excel 文件

    Excel 包含以下 Sheet：
    - Layers: 图层信息
    - Lines: 直线
    - Circles: 圆
    - Arcs: 弧
    - Polyline: 多段线
    - Rectangles: 矩形（从多段线中自动识别）
    - Annotations: 文字标注
    """
    if not file.filename:
        raise HTTPException(400, "缺少文件名")

    lower_name = file.filename.lower()
    if not lower_name.endswith(".dxf"):
        raise HTTPException(400, "只支持 .dxf 文件")

    try:
        import ezdxf  # noqa: F401
    except ImportError as e:
        raise HTTPException(500, "ezdxf 未安装") from e

    try:
        from openpyxl import Workbook  # noqa: F401
    except ImportError as e:
        raise HTTPException(500, "openpyxl 未安装") from e

    # 保存上传文件到临时路径
    import uuid
    job_id = uuid.uuid4().hex
    job_dir = WORK_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    dxf_path = job_dir / file.filename

    try:
        content = await file.read()
        with open(dxf_path, "wb") as f:
            f.write(content)

        excel_bytes = dxf_to_excel_bytes(dxf_path)

        # 生成输出文件名
        output_name = file.filename.rsplit(".", 1)[0] + ".xlsx"

        return StreamingResponse(
            io.BytesIO(excel_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{output_name}"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        LOG.exception("DXF 转 Excel 失败")
        raise HTTPException(500, f"DXF 转 Excel 失败: {e}") from e
    finally:
        import shutil
        shutil.rmtree(job_dir, ignore_errors=True)
